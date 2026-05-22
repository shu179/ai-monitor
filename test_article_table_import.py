import os
import json
from io import BytesIO
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch

import core.article_store as article_store
from backend_lib.article_service import (
    _ARTICLE_IMPORT_BATCHES_LOCK,
    _article_to_api,
    _article_import_batches_lock_file,
    _load_article_import_batches_file,
    _save_article_import_batches_file,
)
from core.file_lock import CrossProcessRLock
from web_backend import (
    ArticleImportBatchStore,
    ArticleService,
    _extract_article_import_items,
    _article_published_date,
    _normalize_article_import_media_type,
    _split_article_import_platform_account,
)


class _ArticleImportConfigProvider:
    def load(self):
        return {
            "tasks": [
                {
                    "name": "品牌A",
                    "brand": "品牌A",
                    "keywords": [{"keyword": "品牌A", "brand": "品牌A"}],
                }
            ]
        }


def _article_import_service(store: ArticleImportBatchStore | None = None) -> tuple[ArticleService, ArticleImportBatchStore]:
    batch_store = store or ArticleImportBatchStore()
    return (
        ArticleService(
            config_provider=_ArticleImportConfigProvider(),
            synced_articles_loader=lambda config: [],
            invalidate_article_cache=lambda: None,
            lock=threading.RLock(),
            import_batch_store=batch_store,
        ),
        batch_store,
    )


class ArticleTableImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self._original_data_dir = os.environ.get("AIBRANDMONITOR_DATA_DIR")
        self._original_article_store_backend = os.environ.get("AIBRANDMONITOR_ARTICLE_STORE_BACKEND")
        os.environ["AIBRANDMONITOR_DATA_DIR"] = self._tmpdir.name
        os.environ["AIBRANDMONITOR_ARTICLE_STORE_BACKEND"] = "json"
        self._original_paths = {
            "ARTICLES_FILE": article_store.ARTICLES_FILE,
            "DOMAIN_OVERRIDES_FILE": article_store.DOMAIN_OVERRIDES_FILE,
            "DOMAIN_MEDIA_NAMES_FILE": article_store.DOMAIN_MEDIA_NAMES_FILE,
            "EXCLUDED_ARTICLE_URLS_FILE": article_store.EXCLUDED_ARTICLE_URLS_FILE,
        }
        root = Path(self._tmpdir.name)
        article_store.ARTICLES_FILE = root / "logs" / "articles.json"
        article_store.DOMAIN_OVERRIDES_FILE = root / "logs" / "domain_overrides.json"
        article_store.DOMAIN_MEDIA_NAMES_FILE = root / "logs" / "domain_media_names.json"
        article_store.EXCLUDED_ARTICLE_URLS_FILE = root / "logs" / "excluded_article_urls.json"
        self._import_batches_path = root / "logs" / "article_import_batches.json"
        self._article_import_batches_path_patch = patch(
            "backend_lib.article_service._article_import_batches_path",
            return_value=self._import_batches_path,
        )
        self._article_import_batches_path_patch.start()

    def tearDown(self) -> None:
        self._article_import_batches_path_patch.stop()
        article_store.ARTICLES_FILE = self._original_paths["ARTICLES_FILE"]
        article_store.DOMAIN_OVERRIDES_FILE = self._original_paths["DOMAIN_OVERRIDES_FILE"]
        article_store.DOMAIN_MEDIA_NAMES_FILE = self._original_paths["DOMAIN_MEDIA_NAMES_FILE"]
        article_store.EXCLUDED_ARTICLE_URLS_FILE = self._original_paths["EXCLUDED_ARTICLE_URLS_FILE"]
        if self._original_data_dir is None:
            os.environ.pop("AIBRANDMONITOR_DATA_DIR", None)
        else:
            os.environ["AIBRANDMONITOR_DATA_DIR"] = self._original_data_dir
        if self._original_article_store_backend is None:
            os.environ.pop("AIBRANDMONITOR_ARTICLE_STORE_BACKEND", None)
        else:
            os.environ["AIBRANDMONITOR_ARTICLE_STORE_BACKEND"] = self._original_article_store_backend
        self._tmpdir.cleanup()

    def _build_workbook(self) -> bytes:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "文章"
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append(["品牌A 发布新品", "https://example.com/news/a", "示例媒体", "2024-01-03"])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    def test_import_table_adds_articles_with_matches_and_can_undo(self) -> None:
        service, batch_store = _article_import_service()

        result = service.import_articles_from_file("文章导入.xlsx", self._build_workbook())

        self.assertTrue(result["ok"])
        self.assertEqual(result["added_count"], 1)
        import_id = result["import_id"]
        articles = article_store.get_articles()
        self.assertEqual(len(articles), 1)
        self.assertEqual(articles[0].get("media_name"), "示例媒体")
        self.assertIn("品牌A", articles[0].get("matched_tasks") or [])
        self.assertIn(import_id, batch_store.get_batches())

        undo_result = service.undo_article_import(import_id)

        self.assertTrue(undo_result["ok"])
        self.assertEqual(article_store.get_articles(), [])
        self.assertFalse(article_store.is_article_url_excluded("https://example.com/news/a"))

    def test_pending_import_batch_survives_runtime_reload(self) -> None:
        service, _ = _article_import_service()

        result = service.import_articles_from_file("文章导入.xlsx", self._build_workbook())
        self.assertTrue(result["ok"])

        reloaded_service, _ = _article_import_service()
        pending = reloaded_service.get_pending_article_imports()

        self.assertTrue(pending["ok"])
        self.assertEqual(len(pending["batches"]), 1)
        self.assertEqual(pending["batches"][0]["import_id"], result["import_id"])

    def test_import_batch_store_merge_uses_cross_process_lock(self) -> None:
        self.assertIsInstance(_ARTICLE_IMPORT_BATCHES_LOCK, CrossProcessRLock)
        _save_article_import_batches_file({
            "existing": {
                "id": "existing",
                "file_name": "existing.xlsx",
                "article_ids": ["article-a"],
                "status": "pending",
            }
        })

        store = ArticleImportBatchStore({
            "new": {
                "id": "new",
                "file_name": "new.xlsx",
                "article_ids": ["article-b"],
                "status": "pending",
            }
        })
        store.save(merge_existing=True)
        loaded = _load_article_import_batches_file()

        self.assertEqual(set(loaded), {"existing", "new"})
        self.assertEqual(
            _article_import_batches_lock_file(),
            self._import_batches_path.with_name(".article_import_batches.json.lock"),
        )
        self.assertTrue(_article_import_batches_lock_file().exists())

    def test_import_media_type_keeps_self_media(self) -> None:
        self.assertEqual(
            _normalize_article_import_media_type("自媒体", "https://example.com/a", "示例媒体"),
            "selfmedia",
        )

    def test_import_detects_side_by_side_article_tables(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "文章汇总"
        worksheet.append([
            "文章标题", "媒体名称", "发布时间", "发布链接", "",
            "文章标题", "自媒体名称", "发布时间", "发布链接",
        ])
        worksheet.append([
            "品牌A 权威报道", "示例媒体", "2024-01-03", "https://example.com/media-a", "",
            "品牌A 自媒体报道", "示例账号", "2024-01-04", "https://example.com/self-a",
        ])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("横排文章.xlsx", output.getvalue())

        self.assertEqual(len(items), 2)
        self.assertEqual([item["title"] for item in items], ["品牌A 权威报道", "品牌A 自媒体报道"])
        self.assertEqual(items[1]["account_name"], "示例账号")
        self.assertEqual(items[1]["media_type"], "selfmedia")
        self.assertEqual([sheet["count"] for sheet in details["sheets"]], [1, 1])

    def test_import_detects_reordered_columns_and_late_header(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "不规则"
        worksheet.append(["客户", "品牌A"])
        worksheet.append(["月份", "2024-01"])
        worksheet.append([])
        worksheet.append(["来源媒体", "发布日期", "文章标题", "原文链接"])
        worksheet.append(["示例媒体", "2024-01-05", "品牌A 不规则表头报道", "https://example.com/reordered"])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("不规则.xlsx", output.getvalue())

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["title"], "品牌A 不规则表头报道")
        self.assertEqual(items[0]["media_name"], "示例媒体")
        self.assertEqual(items[0]["published_at"], "2024-01-05")
        self.assertEqual(details["sheets"][0]["count"], 1)

    def test_import_detects_vertical_article_tables(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "多表"
        worksheet.append(["文章标题", "媒体名称", "发布时间", "发布链接"])
        worksheet.append(["品牌A 上半区报道", "示例媒体", "2024-01-07", "https://example.com/top"])
        worksheet.append([])
        worksheet.append(["文章标题", "自媒体名称", "发布时间", "发布链接"])
        worksheet.append(["品牌A 下半区报道", "头条-示例账号", "2024-01-08", "https://www.toutiao.com/article/7615924630036955694/"])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("纵向多表.xlsx", output.getvalue())

        self.assertEqual([item["title"] for item in items], ["品牌A 上半区报道", "品牌A 下半区报道"])
        self.assertEqual(items[1]["account_name"], "头条-示例账号")
        self.assertEqual([sheet["count"] for sheet in details["sheets"]], [1, 1])

    def test_import_reads_generic_hyperlink_cells(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append(["品牌A 超链接报道", "点击查看", "示例媒体", "2024-01-09"])
        worksheet["B2"].hyperlink = "https://example.com/hyperlink"
        output = BytesIO()
        workbook.save(output)

        items, _ = _extract_article_import_items("超链接.xlsx", output.getvalue())

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["url"], "https://example.com/hyperlink")

    def test_import_uses_title_hyperlink_as_article_url(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "来源媒体", "发布时间"])
        worksheet.append(["品牌A 标题带链接报道", "示例媒体", "2024-01-09"])
        worksheet["A2"].hyperlink = "https://example.com/title-hyperlink"
        output = BytesIO()
        workbook.save(output)

        items, _ = _extract_article_import_items("标题超链接.xlsx", output.getvalue())

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["title"], "品牌A 标题带链接报道")
        self.assertEqual(items[0]["url"], "https://example.com/title-hyperlink")

    def test_import_self_media_name_is_stored_as_account_name(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "自媒体名称", "发布时间", "发布链接"])
        worksheet.append([
            "品牌A 头条账号文章",
            "野渡泛舟客",
            "2024-01-06",
            "https://www.toutiao.com/article/7615924630036955694/",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("自媒体文章.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("media_name"), "今日头条")
        self.assertEqual(article.get("account_name"), "野渡泛舟客")
        self.assertEqual(article.get("media_type"), "selfmedia")

    def test_import_toutiao_alias_uses_full_media_name(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append([
            "品牌A 头条媒体名文章",
            "https://www.toutiao.com/article/7615924630036955694/",
            "头条",
            "2024-01-06",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("头条别名.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("media_name"), "今日头条")

    def test_import_recognizes_publish_url_reject_reason_column(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "发布网址/拒稿理由", "媒体名称", "发布时间"])
        worksheet.append([
            "品牌A 招标采购报道",
            "https://www.gc-zb.com/about/read/id/11471.html",
            "招标与采购网（GEO）",
            "2026-05-08 14:49:37",
        ])
        worksheet.append([
            "品牌A 红商报道",
            "http://www.redsh.com/pinpai/20260508/115118.shtml",
            "红商网（官方）",
            "2026-05-08 12:15:04",
        ])
        worksheet.append([
            "品牌A 咸宁报道",
            "http://www.xnnews.com.cn/zxsd25178/zx/202605/t20260508_4582053.shtml",
            "咸宁新闻网（可发GEO排名）",
            "2026-05-08 11:32:41",
        ])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("发布网址.xlsx", output.getvalue())

        self.assertEqual(len(items), 3)
        self.assertEqual(details["sheets"][0]["columns"]["url"], "发布网址/拒稿理由")
        self.assertEqual(items[0]["url"], "https://www.gc-zb.com/about/read/id/11471.html")
        self.assertEqual(items[1]["url"], "http://www.redsh.com/pinpai/20260508/115118.shtml")
        self.assertEqual(items[2]["url"], "http://www.xnnews.com.cn/zxsd25178/zx/202605/t20260508_4582053.shtml")

    def test_actual_media_order_files_parse_without_missing_urls(self) -> None:
        cases = [
            (Path("/Users/shuao/Desktop/media-order-20260511133947.xlsx"), 76),
            (Path("/Users/shuao/Desktop/媒体订单记录(1).xlsx"), 299),
        ]
        for path, expected_rows in cases:
            with self.subTest(file=path.name):
                if not path.exists():
                    self.skipTest(f"缺少实际测试文件：{path}")
                items, _ = _extract_article_import_items(path.name, path.read_bytes())
                missing_urls = [
                    item for item in items
                    if not str(item.get("url") or "").strip()
                ]

                self.assertEqual(len(items), expected_rows)
                self.assertEqual(missing_urls, [])

    def test_actual_media_order_import_dedupes_same_url_for_json_and_sqlite(self) -> None:
        path = Path("/Users/shuao/Desktop/media-order-20260511133947.xlsx")
        if not path.exists():
            self.skipTest(f"缺少实际测试文件：{path}")
        original_db_file = article_store.ARTICLE_STORE_DB_FILE
        root = Path(self._tmpdir.name)
        try:
            for backend in ("json", "sqlite"):
                with self.subTest(backend=backend):
                    os.environ["AIBRANDMONITOR_ARTICLE_STORE_BACKEND"] = backend
                    article_store.ARTICLE_STORE_DB_FILE = root / "logs" / f"article_store_{backend}.sqlite3"
                    for target in (article_store.ARTICLES_FILE, article_store.ARTICLE_STORE_DB_FILE):
                        try:
                            target.unlink()
                        except FileNotFoundError:
                            pass
                    service, _ = _article_import_service()

                    result = service.import_articles_from_file(path.name, path.read_bytes())
                    articles = article_store.get_articles()
                    smzdm_url = article_store.normalize_article_url("https://post.smzdm.com/zz/p/aqzedvpx/")
                    smzdm_articles = [
                        article for article in articles
                        if article_store.normalize_article_url(
                            article_store.resolve_article_display_url(article)
                        ) == smzdm_url
                    ]

                    self.assertTrue(result["ok"])
                    self.assertEqual(result["added_count"], 70)
                    self.assertEqual(result["duplicate_count"], 6)
                    self.assertEqual(len(articles), 70)
                    self.assertEqual(len(smzdm_articles), 1)
                    self.assertFalse([
                        article for article in articles
                        if not article_store.normalize_article_url(
                            article_store.resolve_article_display_url(article)
                        )
                    ])
        finally:
            article_store.ARTICLE_STORE_DB_FILE = original_db_file

    def test_actual_order_record_import_preserves_redhongan_display_urls(self) -> None:
        path = Path("/Users/shuao/Desktop/媒体订单记录(1).xlsx")
        if not path.exists():
            self.skipTest(f"缺少实际测试文件：{path}")
        original_db_file = article_store.ARTICLE_STORE_DB_FILE
        expected_urls = [
            "https://m.redhongan.com/p/200044.html?timestamp=1778468821795",
            "https://m.redhongan.com/p/200008.html?timestamp=1778467837147",
            "https://www.redhongan.com/p/199231.html",
            "https://m.redhongan.com/p/198758.html?timestamp=1778221397826",
        ]
        root = Path(self._tmpdir.name)
        try:
            for backend in ("json", "sqlite"):
                with self.subTest(backend=backend):
                    os.environ["AIBRANDMONITOR_ARTICLE_STORE_BACKEND"] = backend
                    article_store.ARTICLE_STORE_DB_FILE = root / "logs" / f"article_store_redhongan_{backend}.sqlite3"
                    for target in (article_store.ARTICLES_FILE, article_store.ARTICLE_STORE_DB_FILE):
                        try:
                            target.unlink()
                        except FileNotFoundError:
                            pass
                    service, _ = _article_import_service()

                    result = service.import_articles_from_file(path.name, path.read_bytes())
                    articles = article_store.get_articles()
                    redhongan_urls = [
                        article_store.resolve_article_display_url(article)
                        for article in articles
                        if article.get("media_name") == "红安网"
                    ]

                    self.assertTrue(result["ok"])
                    self.assertEqual(result["added_count"], 299)
                    self.assertEqual(result["duplicate_count"], 0)
                    self.assertCountEqual(redhongan_urls, expected_urls)
        finally:
            article_store.ARTICLE_STORE_DB_FILE = original_db_file

    def test_import_keeps_original_publish_url_for_display(self) -> None:
        from openpyxl import Workbook

        original_url = "https://m.redhongan.com/p/200044.html?timestamp=1778468821795"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "发布链接", "媒体名称", "发布时间"])
        worksheet.append(["品牌A 红安报道", original_url, "红安网", "2026-05-11 13:30:00"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("原始链接.xlsx", output.getvalue())
        stored = article_store.get_articles()[0]

        self.assertTrue(result["ok"])
        self.assertEqual(stored.get("raw_url"), original_url)
        self.assertEqual(stored.get("url"), "https://m.redhongan.com/p/200044.html")
        self.assertEqual(_article_to_api(stored)["url"], original_url)
        self.assertIsNotNone(article_store.find_article_by_url(original_url))

    def test_import_ignores_reject_reason_in_publish_url_column(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "发布网址/拒稿理由", "媒体名称", "发布时间"])
        worksheet.append(["品牌A 被拒稿件", "内容不符合发布要求", "红商网（官方）", "2026-05-08 12:15:04"])
        output = BytesIO()
        workbook.save(output)

        items, _ = _extract_article_import_items("拒稿理由.xlsx", output.getvalue())

        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["url"], "")

    def test_import_extracts_url_from_embedded_and_unmapped_cells(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "媒体名称", "备注", "发布时间", "发布状态"])
        worksheet.append([
            "品牌A 备注里带链接",
            "示例媒体",
            "已发布：https://example.com/from-note?x=1",
            "2026-05-08",
            "",
        ])
        worksheet.append([
            "品牌A 未映射超链接",
            "示例媒体",
            "",
            "2026-05-08",
            "打开",
        ])
        worksheet["E3"].hyperlink = "https://example.com/from-unmapped-hyperlink"
        worksheet.append([
            "品牌A 只有拒稿理由",
            "示例媒体",
            "内容不符合发布要求",
            "2026-05-08",
            "拒稿",
        ])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("未映射链接.xlsx", output.getvalue())

        self.assertEqual(len(items), 3)
        self.assertNotIn("url", details["sheets"][0]["columns"])
        self.assertEqual(items[0]["url"], "https://example.com/from-note?x=1")
        self.assertEqual(items[1]["url"], "https://example.com/from-unmapped-hyperlink")
        self.assertEqual(items[2]["url"], "")

    def test_import_does_not_treat_reject_reason_column_as_url(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "拒稿理由", "媒体名称", "发布时间"])
        worksheet.append(["品牌A 被拒稿件", "内容不符合发布要求", "红商网（官方）", "2026-05-08 12:15:04"])
        output = BytesIO()
        workbook.save(output)

        items, details = _extract_article_import_items("拒稿理由列.xlsx", output.getvalue())

        self.assertEqual(len(items), 1)
        self.assertNotIn("url", details["sheets"][0]["columns"])
        self.assertEqual(items[0]["url"], "")

    def test_import_tencent_news_bracket_media_uses_account_name(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "发布网址/拒稿理由", "媒体名称", "发布时间"])
        worksheet.append([
            "品牌A 腾讯新闻账号文章",
            "https://page.om.qq.com/page/OLbKqsxlvaA6jHk3rr7F86DA0",
            "无线昆明（腾讯新闻）",
            "2026-04-30 17:50:05",
        ])
        worksheet.append([
            "品牌A 官方腾讯号文章",
            "https://view.inews.qq.com/a/20260428A03XKS00",
            "濮阳市广播电视台（官方腾讯号）",
            "2026-04-28 11:50:05",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("腾讯新闻账号.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        articles = {article.get("title"): article for article in article_store.get_articles()}
        wireless = articles["品牌A 腾讯新闻账号文章"]
        self.assertEqual(wireless.get("media_name"), "腾讯新闻")
        self.assertEqual(wireless.get("account_name"), "无线昆明")
        self.assertEqual(wireless.get("media_type"), "selfmedia")
        official = articles["品牌A 官方腾讯号文章"]
        self.assertEqual(official.get("media_name"), "腾讯新闻")
        self.assertEqual(official.get("account_name"), "濮阳市广播电视台")
        self.assertEqual(official.get("media_type"), "selfmedia")

    def test_import_strips_order_qualifier_without_creating_account_name(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append([
            "品牌A 红商报道",
            "http://www.redsh.com/pinpai/20260508/115118.shtml",
            "红商网（官方）",
            "2026-05-08 12:15:04",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("媒体标签.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("media_name"), "红商网")
        self.assertFalse(article.get("account_name"))
        self.assertEqual(article.get("media_type"), "authority")

    def test_self_media_platform_domain_repairs_learned_account_name(self) -> None:
        article_store.DOMAIN_MEDIA_NAMES_FILE.parent.mkdir(parents=True, exist_ok=True)
        article_store.DOMAIN_MEDIA_NAMES_FILE.write_text(
            json.dumps({"sohu.com": "时尚潮流家"}, ensure_ascii=False),
            encoding="utf-8",
        )
        article_store.ARTICLES_FILE.parent.mkdir(parents=True, exist_ok=True)
        article_store.ARTICLES_FILE.write_text(
            json.dumps(
                [
                    {
                        "id": "dirty-sohu",
                        "url": "https://sohu.com/a/995634977_120983381",
                        "title": "品牌A 搜狐文章",
                        "media_name": "时尚潮流家",
                        "media_type": "selfmedia",
                        "ts": "2024-01-06",
                    }
                ],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        article = article_store.get_articles()[0]
        stored = json.loads(article_store.ARTICLES_FILE.read_text(encoding="utf-8"))[0]

        self.assertEqual(article.get("media_name"), "搜狐")
        self.assertEqual(article.get("account_name"), "时尚潮流家")
        self.assertEqual(article_store.resolve_article_source(article), "搜狐")
        self.assertEqual(stored.get("media_name"), "搜狐")
        self.assertEqual(stored.get("account_name"), "时尚潮流家")

        article_store.save_domain_media_name("sohu.com", "另一个搜狐账号", force=True)
        media_names = json.loads(article_store.DOMAIN_MEDIA_NAMES_FILE.read_text(encoding="utf-8"))
        self.assertEqual(media_names.get("sohu.com"), "搜狐")

    def test_legacy_article_records_repair_media_and_account_on_load(self) -> None:
        article_store.ARTICLES_FILE.parent.mkdir(parents=True, exist_ok=True)
        article_store.ARTICLES_FILE.write_text(
            json.dumps(
                [
                    {
                        "id": "legacy-tencent-bracket",
                        "url": "https://page.om.qq.com/page/OLbKqsxlvaA6jHk3rr7F86DA0",
                        "title": "品牌A 腾讯新闻账号文章",
                        "media_name": "无线昆明（腾讯新闻）",
                        "media_type": "authority",
                        "ts": "2026-04-30",
                    },
                    {
                        "id": "legacy-tencent-swapped",
                        "url": "https://view.inews.qq.com/a/20260428A03XKS00",
                        "title": "品牌A 官方腾讯号文章",
                        "media_name": "濮阳市广播电视台",
                        "account_name": "官方腾讯号",
                        "media_type": "authority",
                        "ts": "2026-04-28",
                    },
                    {
                        "id": "legacy-redsh",
                        "url": "http://www.redsh.com/pinpai/20260508/115118.shtml",
                        "title": "品牌A 红商报道",
                        "media_name": "红商网（官方）",
                        "media_type": "selfmedia",
                        "ts": "2026-05-08",
                    },
                    {
                        "id": "legacy-gczb",
                        "url": "https://www.gc-zb.com/about/read/id/11471.html",
                        "title": "品牌A 招标采购报道",
                        "media_name": "招标与采购网（GEO）",
                        "media_type": "selfmedia",
                        "ts": "2026-05-08",
                    },
                ],
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )

        articles = {article["id"]: article for article in article_store.get_articles()}
        stored = {
            article["id"]: article
            for article in json.loads(article_store.ARTICLES_FILE.read_text(encoding="utf-8"))
        }

        self.assertEqual(articles["legacy-tencent-bracket"].get("media_name"), "腾讯新闻")
        self.assertEqual(articles["legacy-tencent-bracket"].get("account_name"), "无线昆明")
        self.assertEqual(articles["legacy-tencent-bracket"].get("media_type"), "selfmedia")
        self.assertEqual(articles["legacy-tencent-swapped"].get("media_name"), "腾讯新闻")
        self.assertEqual(articles["legacy-tencent-swapped"].get("account_name"), "濮阳市广播电视台")
        self.assertEqual(articles["legacy-tencent-swapped"].get("media_type"), "selfmedia")
        self.assertEqual(articles["legacy-redsh"].get("media_name"), "红商网")
        self.assertEqual(articles["legacy-redsh"].get("media_type"), "authority")
        self.assertEqual(articles["legacy-gczb"].get("media_name"), "招标与采购网")
        self.assertEqual(articles["legacy-gczb"].get("media_type"), "authority")
        self.assertEqual(stored["legacy-tencent-bracket"].get("account_name"), "无线昆明")
        self.assertEqual(stored["legacy-redsh"].get("media_name"), "红商网")

    def test_import_self_media_account_name_in_media_column_is_kept_as_account_name(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "媒体名称", "发布时间", "发布链接"])
        worksheet.append([
            "品牌A 搜狐账号文章",
            "品牌测评中心",
            "2024-01-06",
            "https://sohu.com/a/995634977_120983381",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("搜狐账号列.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("media_name"), "搜狐")
        self.assertEqual(article.get("account_name"), "品牌测评中心")

    def test_import_without_publish_date_stays_undated(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体"])
        worksheet.append(["品牌A 缺少日期报道", "https://example.com/no-date", "示例媒体"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("缺日期.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("published_at"), "")
        self.assertEqual(article.get("ts"), "")
        self.assertIsNone(_article_published_date(article))

    def test_import_existing_url_updates_metadata_and_undo_restores(self) -> None:
        original = article_store.add_article({
            "url": "https://example.com/existing",
            "title": "旧标题",
            "media_name": "旧媒体",
            "media_type": "authority",
            "published_at": "2024-01-01",
            "ts": "2024-01-01",
            "matched_tasks": [],
            "match_reasons": {},
            "unmatched_reason": "",
        })
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append(["品牌A 更新后的标题", "https://example.com/existing", "新媒体", "2024-01-10"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("品牌A更新.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        self.assertEqual(result["added_count"], 0)
        self.assertEqual(result["updated_count"], 1)
        updated = article_store.find_article_by_url("https://example.com/existing")
        self.assertEqual(updated.get("title"), "品牌A 更新后的标题")
        self.assertEqual(updated.get("published_at"), "2024-01-10")
        self.assertIn("品牌A", updated.get("matched_tasks") or [])

        undo_result = service.undo_article_import(result["import_id"])

        self.assertTrue(undo_result["ok"])
        restored = article_store.find_article_by_url("https://example.com/existing")
        self.assertEqual(restored.get("id"), original.get("id"))
        self.assertEqual(restored.get("title"), "旧标题")
        self.assertEqual(restored.get("media_name"), "旧媒体")
        self.assertEqual(restored.get("published_at"), "2024-01-01")

    def test_import_repairs_existing_missing_url_record(self) -> None:
        original = article_store.add_article({
            "id": "missing-url",
            "url": "",
            "title": "品牌A 红安报道",
            "media_name": "红安网",
            "media_type": "authority",
            "published_at": "2026-05-11",
            "ts": "2026-05-11",
            "matched_tasks": ["品牌A"],
        })
        from openpyxl import Workbook

        original_url = "https://m.redhongan.com/p/200044.html?timestamp=1778468821795"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "发布链接", "媒体名称", "发布时间"])
        worksheet.append(["品牌A 红安报道", original_url, "红安网", "2026-05-11 13:30:00"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("修复缺链接.xlsx", output.getvalue())
        articles = article_store.get_articles()

        self.assertTrue(result["ok"])
        self.assertEqual(result["added_count"], 0)
        self.assertEqual(result["updated_count"], 1)
        self.assertEqual(len(articles), 1)
        self.assertEqual(articles[0].get("id"), original.get("id"))
        self.assertEqual(articles[0].get("url"), "https://m.redhongan.com/p/200044.html")
        self.assertEqual(articles[0].get("raw_url"), original_url)

    def test_import_existing_url_preserves_deleted_task_classification(self) -> None:
        article_store.add_article({
            "url": "https://example.com/deleted-task",
            "title": "已删除品牌历史文章",
            "media_name": "旧媒体",
            "media_type": "authority",
            "published_at": "2024-01-01",
            "ts": "2024-01-01",
            "matched_tasks": ["已删除品牌"],
            "match_reasons": {"已删除品牌": ["历史归类"]},
            "unmatched_reason": "",
        })
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append(["行业更新报道", "https://example.com/deleted-task", "新媒体", "2024-01-10"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("历史任务更新.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        self.assertEqual(result["updated_count"], 1)
        updated = article_store.get_articles()[0]
        self.assertEqual(updated.get("matched_tasks"), ["已删除品牌"])
        self.assertEqual(updated.get("match_reasons"), {"已删除品牌": ["历史归类"]})

    def test_update_article_preserves_deleted_task_classification(self) -> None:
        article = article_store.add_article({
            "url": "https://example.com/edit-deleted-task",
            "title": "已删除品牌历史文章",
            "media_name": "旧媒体",
            "media_type": "authority",
            "published_at": "2024-01-01",
            "ts": "2024-01-01",
            "matched_tasks": ["已删除品牌"],
            "match_reasons": {"已删除品牌": ["历史归类"]},
            "unmatched_reason": "",
        })
        service, _ = _article_import_service()

        result = service.update_article(article["id"], {"title": "行业更新报道"})

        self.assertTrue(result["ok"])
        updated = article_store.get_articles()[0]
        self.assertEqual(updated.get("title"), "行业更新报道")
        self.assertEqual(updated.get("matched_tasks"), ["已删除品牌"])
        self.assertEqual(updated.get("match_reasons"), {"已删除品牌": ["历史归类"]})

    def test_import_splits_platform_account_text(self) -> None:
        self.assertEqual(_split_article_import_platform_account("头条（野渡泛舟客）"), ("今日头条", "野渡泛舟客"))
        self.assertEqual(_split_article_import_platform_account("头条-野渡泛舟客"), ("今日头条", "野渡泛舟客"))
        self.assertEqual(_split_article_import_platform_account("搜狐号/品牌测评榜单"), ("搜狐号", "品牌测评榜单"))
        self.assertEqual(_split_article_import_platform_account("36氪-出海"), ("", ""))

    def test_import_account_column_can_include_platform_account_text(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "自媒体名称", "发布时间", "发布链接"])
        worksheet.append([
            "品牌A 组合账号文章",
            "头条（野渡泛舟客）",
            "2024-01-06",
            "https://www.toutiao.com/article/7615924630036955694/",
        ])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("组合账号.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertEqual(article.get("media_name"), "今日头条")
        self.assertEqual(article.get("account_name"), "野渡泛舟客")

    def test_import_uses_file_name_as_classification_hint(self) -> None:
        from openpyxl import Workbook

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["文章标题", "文章链接", "来源媒体", "发布时间"])
        worksheet.append(["没有品牌词的行业报道", "https://example.com/file-hint", "示例媒体", "2024-01-11"])
        output = BytesIO()
        workbook.save(output)
        service, _ = _article_import_service()

        result = service.import_articles_from_file("GEO-品牌A-3月汇总.xlsx", output.getvalue())

        self.assertTrue(result["ok"])
        article = article_store.get_articles()[0]
        self.assertIn("品牌A", article.get("matched_tasks") or [])


if __name__ == "__main__":
    unittest.main()
