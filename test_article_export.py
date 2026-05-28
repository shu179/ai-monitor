from openpyxl import load_workbook

from core.article_store import (
    build_article_export_keyword_cache,
    build_article_export_keyword_signature,
    resolve_article_export_keywords,
)
from backend_lib.article_export import _article_export_items as _backend_article_export_items
from web_backend import _build_article_export_xlsx


def test_article_export_xlsx_includes_keyword_category_and_account(tmp_path):
    output_path = tmp_path / "articles.xlsx"
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [
                    {"keyword": "融资"},
                    {"keyword": "新品"},
                ],
            }
        ]
    }
    articles = [
        {
            "id": "selfmedia",
            "title": "品牌A 新品 发布",
            "url": "https://example.com/self",
            "media_name": "今日头条",
            "media_type": "selfmedia",
            "account_name": "作者甲",
            "published_at": "2024-01-02",
            "matched_tasks": ["品牌A"],
        },
        {
            "id": "authority",
            "title": "品牌A 融资 新闻",
            "url": "https://example.com/news",
            "media_name": "财新",
            "media_type": "authority",
            "published_at": "2024/01/01 09:30",
            "matched_tasks": ["品牌A"],
        },
    ]

    _build_article_export_xlsx(
        brand_name="品牌A",
        start_date="2024-01-01",
        end_date="2024-01-31",
        articles=articles,
        output_path=output_path,
        show_keyword_category=True,
        show_selfmedia_account=True,
        config=config,
        task_name="品牌A",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active

    assert sheet["A1"].value == "品牌A 文章汇总"
    assert [sheet.cell(row=5, column=col).value for col in range(1, 7)] == [
        "关键词大类",
        "来源",
        "类型",
        "文章标题",
        "文章链接",
        "发布时间",
    ]
    assert [sheet.cell(row=6, column=col).value for col in range(1, 7)] == [
        "融资",
        "财新",
        "权威媒体",
        "品牌A 融资 新闻",
        "https://example.com/news",
        "2024-01-01",
    ]
    assert [sheet.cell(row=7, column=col).value for col in range(1, 7)] == [
        "新品",
        "今日头条（作者甲）",
        "自媒体",
        "品牌A 新品 发布",
        "https://example.com/self",
        "2024-01-02",
    ]


def test_article_export_keyword_category_keeps_global_time_order(tmp_path):
    output_path = tmp_path / "articles.xlsx"
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "新品"}, {"keyword": "融资"}],
            }
        ]
    }
    articles = [
        {
            "id": "newer-keyword-first",
            "title": "品牌A 新品 发布",
            "url": "https://example.com/newer",
            "media_name": "示例媒体",
            "media_type": "authority",
            "published_at": "2024-01-02",
            "matched_tasks": ["品牌A"],
        },
        {
            "id": "older-keyword-second",
            "title": "品牌A 融资 新闻",
            "url": "https://example.com/older",
            "media_name": "示例媒体",
            "media_type": "authority",
            "published_at": "2024-01-01",
            "matched_tasks": ["品牌A"],
        },
    ]

    _build_article_export_xlsx(
        brand_name="品牌A",
        start_date="",
        end_date="",
        articles=articles,
        output_path=output_path,
        show_keyword_category=True,
        show_selfmedia_account=True,
        config=config,
        task_name="品牌A",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active

    assert sheet.cell(row=6, column=1).value == "融资"
    assert sheet.cell(row=6, column=4).value == "品牌A 融资 新闻"
    assert sheet.cell(row=7, column=1).value == "新品"
    assert sheet.cell(row=7, column=4).value == "品牌A 新品 发布"


def test_article_export_keyword_category_is_one_cell_not_row_grouping():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "融资"}, {"keyword": "新品"}],
            }
        ]
    }
    article = {
        "id": "multi-keyword",
        "title": "品牌A 融资 新品 发布",
        "url": "https://example.com/multi",
        "media_name": "示例媒体",
        "media_type": "authority",
        "published_at": "2024-01-01",
        "matched_tasks": ["品牌A"],
    }

    assert _backend_article_export_items(
        [article],
        show_keyword_category=True,
        config=config,
        task_name="品牌A",
    ) == [("融资 · 新品", article)]


def test_article_export_keyword_category_joiner_isolates_keywords_with_inner_punctuation():
    """关键词内部带顿号/逗号也不能和"多关键词拼接分隔符"混淆，
    历史上拼接用"、"导致 '保湿、防脱' 看着像一个怪关键词；现在改成 ' · '。"""
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [
                    {"keyword": "干皮水润保湿、服帖不卡粉的粉底液推荐"},
                    {"keyword": "敏感肌可用"},
                ],
            }
        ]
    }
    article = {
        "id": "punctuation",
        "title": "干皮水润保湿、服帖不卡粉的粉底液推荐 敏感肌可用",
        "media_name": "示例媒体",
        "media_type": "authority",
        "published_at": "2024-01-01",
        "matched_tasks": ["品牌A"],
    }
    items = _backend_article_export_items(
        [article],
        show_keyword_category=True,
        config=config,
        task_name="品牌A",
    )
    assert len(items) == 1
    cell_value, _ = items[0]
    assert " · " in cell_value
    assert "干皮水润保湿、服帖不卡粉的粉底液推荐" in cell_value
    assert "敏感肌可用" in cell_value


def test_desktop_article_export_keyword_category_keeps_global_time_order():
    from ui.article_window import _article_export_items as desktop_article_export_items

    rows = [
        {
            "id": "newer-keyword-first",
            "title": "品牌A 新品 发布",
            "keyword_categories": ["新品"],
            "published_at": "2024-01-02",
        },
        {
            "id": "older-keyword-second",
            "title": "品牌A 融资 新闻",
            "keyword_categories": ["融资"],
            "published_at": "2024-01-01",
        },
    ]

    items = desktop_article_export_items(
        rows,
        show_keyword_category=True,
        keyword_order={"新品": 0, "融资": 1},
    )

    assert [(label, row["id"]) for label, row in items] == [
        ("融资", "older-keyword-second"),
        ("新品", "newer-keyword-first"),
    ]


def test_article_export_xlsx_fills_missing_duplicate_link(tmp_path):
    output_path = tmp_path / "articles.xlsx"
    articles = [
        {
            "id": "without-link",
            "title": "品牌A 同标题报道",
            "url": "",
            "media_name": "示例媒体",
            "media_type": "authority",
            "published_at": "2024-01-02",
            "matched_tasks": ["品牌A"],
        },
        {
            "id": "with-link",
            "title": "品牌A 同标题报道",
            "url": "https://example.com/linked",
            "media_name": "示例媒体",
            "media_type": "authority",
            "published_at": "2024-01-02",
            "matched_tasks": ["品牌A"],
        },
    ]

    _build_article_export_xlsx(
        brand_name="品牌A",
        start_date="",
        end_date="",
        articles=articles,
        output_path=output_path,
        show_keyword_category=False,
        show_selfmedia_account=True,
        config={"tasks": []},
        task_name="品牌A",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active

    assert sheet.cell(row=6, column=4).value == "https://example.com/linked"
    assert sheet.cell(row=7, column=4).value == "https://example.com/linked"


def test_article_export_xlsx_uses_original_import_url(tmp_path):
    output_path = tmp_path / "articles.xlsx"
    original_url = "https://m.redhongan.com/p/200044.html?timestamp=1778468821795"
    articles = [
        {
            "id": "redhongan",
            "title": "品牌A 红安报道",
            "url": "https://m.redhongan.com/p/200044.html",
            "raw_url": original_url,
            "media_name": "红安网",
            "media_type": "authority",
            "published_at": "2026-05-11",
            "matched_tasks": ["品牌A"],
        },
    ]

    _build_article_export_xlsx(
        brand_name="品牌A",
        start_date="",
        end_date="",
        articles=articles,
        output_path=output_path,
        show_keyword_category=False,
        show_selfmedia_account=True,
        config={"tasks": []},
        task_name="品牌A",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active

    assert sheet.cell(row=6, column=4).value == original_url


def test_article_export_keywords_only_use_configured_keywords():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "品牌A 新品"}],
            }
        ]
    }
    article = {
        "title": "品牌A 新闻动态",
        "matched_tasks": ["品牌A"],
    }

    assert resolve_article_export_keywords(article, config, "品牌A") == []


def test_article_export_keywords_do_not_use_inferred_core_terms():
    config = {
        "tasks": [
            {
                "name": "万通",
                "brand": "万通",
                "keywords": [{"keyword": "万通职业教育", "brand": "万通"}],
            }
        ]
    }
    article = {
        "title": "技能驱动未来：安徽职业教育的实践图景",
        "matched_tasks": ["万通"],
    }

    assert resolve_article_export_keywords(article, config, "万通") == []


def test_article_export_keywords_require_full_configured_keyword():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "AI 编程"}],
            }
        ]
    }
    article = {
        "title": "品牌A AI智能体开发与编程选型",
        "matched_tasks": ["品牌A"],
    }

    assert resolve_article_export_keywords(article, config, "品牌A") == []


def test_article_export_short_english_keywords_use_word_boundaries():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "AI"}, {"keyword": "GEO"}],
            }
        ]
    }

    assert resolve_article_export_keywords(
        {"title": "品牌A AI智能体开发", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["AI"]
    assert resolve_article_export_keywords(
        {"title": "品牌A OpenAI智能体开发", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == []
    assert resolve_article_export_keywords(
        {"title": "品牌A GEO服务商测评", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["GEO"]


def test_article_export_keyword_cache_is_used_and_invalidated():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "新品"}],
            }
        ]
    }
    article = {
        "title": "品牌A 新品发布",
        "matched_tasks": ["品牌A"],
    }
    article.update(build_article_export_keyword_cache(article, config))
    signature = build_article_export_keyword_signature(config)

    assert resolve_article_export_keywords(
        article,
        config,
        "品牌A",
        keyword_config_signature=signature,
    ) == ["新品"]

    stale_article = dict(article)
    stale_article["title"] = "品牌A 普通报道"
    assert resolve_article_export_keywords(
        stale_article,
        config,
        "品牌A",
        keyword_config_signature=signature,
    ) == []


def test_article_export_xlsx_falls_back_to_unknown_when_only_brand_matched(tmp_path):
    output_path = tmp_path / "articles.xlsx"
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "品牌A 新品"}],
            }
        ]
    }
    articles = [
        {
            "id": "brand-only",
            "title": "品牌A 新闻动态",
            "url": "https://example.com/brand-only",
            "media_name": "示例媒体",
            "media_type": "authority",
            "published_at": "2024-01-02",
            "matched_tasks": ["品牌A"],
        }
    ]

    _build_article_export_xlsx(
        brand_name="品牌A",
        start_date="",
        end_date="",
        articles=articles,
        output_path=output_path,
        show_keyword_category=True,
        show_selfmedia_account=True,
        config=config,
        task_name="品牌A",
    )

    workbook = load_workbook(output_path, read_only=True, data_only=True)
    sheet = workbook.active

    assert sheet.cell(row=6, column=1).value == "未知"
    assert sheet.cell(row=6, column=4).value == "品牌A 新闻动态"


def test_article_export_keyword_category_fallback_unknown_in_export_items():
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "新品"}],
            }
        ]
    }
    article = {
        "id": "brand-only",
        "title": "品牌A 普通动态",
        "matched_tasks": ["品牌A"],
    }

    items = _backend_article_export_items(
        [article],
        show_keyword_category=True,
        config=config,
        task_name="品牌A",
    )
    assert items == [("未知", article)]


def test_article_export_keywords_match_full_width_letters_and_digits():
    """全角字母/数字标题应能匹配半角关键词。"""
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [
                    {"keyword": "GEO"},
                    {"keyword": "AI 编程"},
                    {"keyword": "GPT-4"},
                ],
            }
        ]
    }

    # 全角字母
    assert resolve_article_export_keywords(
        {"title": "品牌A ＧＥＯ优化方案", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["GEO"]

    # 全角字母 + 全角空格
    assert resolve_article_export_keywords(
        {"title": "品牌A ＡＩ　编程实战", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["AI 编程"]

    # 全角数字 + 半角字母混合
    assert resolve_article_export_keywords(
        {"title": "品牌A GPT-４ 全面解析", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["GPT-4"]


def test_article_export_keywords_tolerate_inserted_filler_words():
    """标题在关键词中间多一两个虚词时仍应命中。"""
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [
                    {"keyword": "数字化转型"},
                    {"keyword": "AI 编程"},
                ],
            }
        ]
    }

    assert resolve_article_export_keywords(
        {"title": "数字化的转型实践图景", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["数字化转型"]

    assert resolve_article_export_keywords(
        {"title": "企业数字化与转型路径", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["数字化转型"]

    assert resolve_article_export_keywords(
        {"title": "AI 优质 编程实战", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == ["AI 编程"]


def test_article_export_keywords_reject_semantically_different_titles():
    """中间塞入不在白名单的内容，必须保持不命中，避免“AI编程→AI教程”这种错配。"""
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [
                    {"keyword": "AI 编程"},
                    {"keyword": "数字化转型"},
                ],
            }
        ]
    }

    # 主体不同的“AI 教程”不应命中“AI 编程”
    assert resolve_article_export_keywords(
        {"title": "AI 教程合集", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == []

    # 中间塞 5 个字以上的修饰词不应命中
    assert resolve_article_export_keywords(
        {"title": "数字化非常重要的转型", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == []

    # 中间塞数字也不命中（数字不是白名单）
    assert resolve_article_export_keywords(
        {"title": "数字化123转型", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == []


def test_article_export_keywords_short_ascii_still_strict_with_boundary():
    """短 ASCII 关键词 (AI/GEO) 仍然走词边界，近似匹配不应放宽它。"""
    config = {
        "tasks": [
            {
                "name": "品牌A",
                "brand": "品牌A",
                "keywords": [{"keyword": "AI"}],
            }
        ]
    }
    assert resolve_article_export_keywords(
        {"title": "品牌A OpenAI 教程", "matched_tasks": ["品牌A"]},
        config,
        "品牌A",
    ) == []


def test_desktop_article_export_items_falls_back_to_unknown():
    """桌面端导出当 keyword_categories 为空时也要回填“未知”。"""
    from ui.article_window import _article_export_items as desktop_article_export_items

    rows = [
        {
            "id": "no-keyword",
            "title": "品牌A 普通动态",
            "keyword_categories": [],
            "published_at": "2024-01-02",
        },
        {
            "id": "with-keyword",
            "title": "品牌A 新品发布",
            "keyword_categories": ["新品"],
            "published_at": "2024-01-01",
        },
    ]
    items = desktop_article_export_items(
        rows,
        show_keyword_category=True,
        keyword_order={"新品": 0},
    )
    assert [(label, row["id"]) for label, row in items] == [
        ("新品", "with-keyword"),
        ("未知", "no-keyword"),
    ]
