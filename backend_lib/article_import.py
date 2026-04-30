"""Article table import parsing helpers for the local web backend."""

from __future__ import annotations

import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from core.time_utils import parse_local_date


_ARTICLE_IMPORT_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "title": ("标题", "文章标题", "新闻标题", "稿件标题", "内容标题", "题名", "题目", "标题名称", "title", "headline", "subject"),
    "url": ("链接", "文章链接", "原文链接", "发布链接", "网址", "地址", "url", "link", "href"),
    "media_name": ("媒体", "媒体名", "媒体名称", "来源", "来源媒体", "发布媒体", "发布平台", "平台", "平台名称", "站点", "站点名称", "source", "media", "outlet", "publisher"),
    "media_type": ("类型", "媒体类型", "来源类型", "平台类型", "分类", "类别", "type", "category"),
    "published_at": ("发布时间", "发布日期", "发表时间", "发文时间", "刊发时间", "推送时间", "日期", "时间", "published", "published_at", "publishdate", "date", "time", "ts"),
    "excerpt": ("摘要", "简介", "导语", "内容", "正文", "摘录", "备注", "summary", "excerpt", "description", "content", "note"),
    "account_name": ("账号", "账号名", "账号名称", "作者", "作者名", "博主", "达人", "自媒体", "自媒体名", "自媒体名称", "account", "author", "user", "creator"),
    "brand_name": ("品牌", "品牌名", "品牌名称", "客户", "客户名", "客户名称", "公司", "公司名", "公司名称", "企业", "企业名称", "brand", "client", "company"),
    "task_name": ("任务", "任务名", "任务名称", "项目", "项目名", "项目名称", "监测任务", "归属任务", "task", "project"),
}


def _normalize_import_header(value: Any) -> str:
    return re.sub(r"[\s_:\-—–|｜/\\（）()【】\\[\\]\"'“”‘’]+", "", str(value or "").strip().lower())


def _match_article_import_field(value: Any) -> str:
    text = str(value or "").strip()
    normalized = _normalize_import_header(text)
    if not normalized:
        return ""
    if len(normalized) > 40 or re.search(r"https?://|www\.|/", text, re.IGNORECASE):
        return ""
    for field, aliases in _ARTICLE_IMPORT_FIELD_ALIASES.items():
        for alias in aliases:
            alias_key = _normalize_import_header(alias)
            if normalized == alias_key:
                return field
    if (
        any(token in normalized for token in ("标题", "题名", "题目", "headline"))
        and len(normalized) <= 8
        and "的" not in normalized
        and (normalized.startswith(("文章", "新闻", "稿件", "内容", "原文", "主")) or normalized in {"标题", "题名", "题目"})
    ):
        return "title"
    if any(token in normalized for token in ("链接", "网址", "地址", "url", "href", "link")):
        return "url"
    if any(token in normalized for token in ("发布时间", "发布日期", "发表时间", "发文时间", "刊发时间", "推送时间", "publishdate", "publishedat")):
        return "published_at"
    if normalized.endswith("时间") or normalized.endswith("日期"):
        return "published_at"
    if any(token in normalized for token in ("自媒体", "账号", "作者", "博主", "达人", "creator")):
        return "account_name"
    if any(token in normalized for token in ("品牌", "客户", "公司", "企业", "brand", "client", "company")):
        return "brand_name"
    if any(token in normalized for token in ("任务", "项目", "监测任务", "归属任务", "task", "project")):
        return "task_name"
    if (
        any(token in normalized for token in ("媒体", "来源", "发布平台", "平台名称", "站点", "publisher", "outlet"))
        and not normalized.startswith("新媒体")
    ):
        return "media_name"
    if "类型" in normalized or normalized in {"分类", "类别", "type", "category"}:
        return "media_type"
    if any(token in normalized for token in ("摘要", "导语", "简介", "正文", "内容", "摘录", "备注", "summary", "excerpt", "description")):
        return "excerpt"
    return ""


def _stringify_table_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _decode_csv_rows(data: bytes) -> list[list[str]]:
    last_error = ""
    text = ""
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = data.decode(encoding)
            break
        except Exception as exc:
            last_error = str(exc)
    if not text:
        raise ValueError(f"无法识别 CSV 编码{f'：{last_error}' if last_error else ''}")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;，")
    except Exception:
        dialect = csv.excel
    return [
        [_stringify_table_cell(cell) for cell in row]
        for row in csv.reader(StringIO(text), dialect)
    ]


def _decode_xlsx_sheets(data: bytes) -> list[tuple[str, list[list[str]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("当前运行环境缺少 openpyxl，无法读取 Excel 表格") from exc

    workbook = load_workbook(BytesIO(data), read_only=False, data_only=True)
    sheets: list[tuple[str, list[list[str]]]] = []

    def cell_text(cell: Any) -> str:
        value = _stringify_table_cell(getattr(cell, "value", None))
        hyperlink = getattr(getattr(cell, "hyperlink", None), "target", "") or ""
        hyperlink = str(hyperlink or "").strip()
        if hyperlink and (
            not value
            or re.fullmatch(r"(点击)?查看|链接|原文|阅读原文|打开|详情|跳转", value, re.IGNORECASE)
        ):
            return hyperlink
        return value

    for worksheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in worksheet.iter_rows():
            values = [cell_text(cell) for cell in row]
            while values and not values[-1]:
                values.pop()
            if any(values):
                rows.append(values)
        if rows:
            sheets.append((worksheet.title, rows))
    return sheets


def _read_article_import_tables(file_name: str, data: bytes) -> list[tuple[str, list[list[str]]]]:
    suffix = Path(str(file_name or "")).suffix.lower()
    if suffix == ".csv":
        return [(Path(file_name).stem or "CSV", _decode_csv_rows(data))]
    return _decode_xlsx_sheets(data)


def _detect_article_import_header(rows: list[list[str]]) -> tuple[int, dict[str, int], int]:
    best_index = -1
    best_score = 0
    best_map: dict[str, int] = {}
    for index, row in enumerate(rows[:8]):
        field_map: dict[str, int] = {}
        score = 0
        for col_index, cell in enumerate(row):
            field = _match_article_import_field(cell)
            if not field or field in field_map:
                continue
            field_map[field] = col_index
            score += 2
        if "title" in field_map:
            score += 4
        if "url" in field_map:
            score += 3
        if "published_at" in field_map:
            score += 2
        if score > best_score:
            best_index = index
            best_score = score
            best_map = field_map
    return best_index, best_map, best_score


def _score_article_import_field_map(field_map: dict[str, int]) -> int:
    score = len(field_map) * 2
    if "title" in field_map:
        score += 4
    if "url" in field_map:
        score += 3
    if "published_at" in field_map:
        score += 2
    return score


def _article_import_header_segments(row: list[str]) -> list[tuple[int, int]]:
    segments: list[tuple[int, int]] = []
    start: int | None = None
    last_non_blank = -1
    for index, cell in enumerate(row):
        if str(cell or "").strip():
            if start is None:
                start = index
            last_non_blank = index
            continue
        if start is not None:
            segments.append((start, last_non_blank))
            start = None
            last_non_blank = -1
    if start is not None:
        segments.append((start, last_non_blank))
    return segments


def _infer_article_import_block_media_type(row: list[str], start_col: int, end_col: int) -> str:
    header_text = "".join(str(cell or "") for cell in row[start_col:end_col + 1])
    if re.search(r"自媒体|账号|博主|达人|creator|account", header_text, re.IGNORECASE):
        return "selfmedia"
    if re.search(r"权威|官媒|官方媒体", header_text, re.IGNORECASE):
        return "authority"
    return ""


def _build_article_import_block(row: list[str], header_index: int, start_col: int, end_col: int) -> dict[str, Any] | None:
    field_map: dict[str, int] = {}
    for col_index in range(start_col, min(end_col + 1, len(row))):
        field = _match_article_import_field(row[col_index])
        if not field or field in field_map:
            continue
        field_map[field] = col_index
    score = _score_article_import_field_map(field_map)
    if "title" not in field_map or score < 6:
        return None
    return {
        "header_index": header_index,
        "field_map": field_map,
        "score": score,
        "start_col": start_col,
        "end_col": end_col,
        "default_media_type": _infer_article_import_block_media_type(row, start_col, end_col),
    }


def _split_article_import_segment_by_titles(
    row: list[str],
    header_index: int,
    start_col: int,
    end_col: int,
) -> list[dict[str, Any]]:
    title_columns = [
        col_index
        for col_index in range(start_col, min(end_col + 1, len(row)))
        if _match_article_import_field(row[col_index]) == "title"
    ]
    if len(title_columns) <= 1:
        block = _build_article_import_block(row, header_index, start_col, end_col)
        return [block] if block else []
    blocks: list[dict[str, Any]] = []
    for index, title_col in enumerate(title_columns):
        block_start = start_col if index == 0 else title_col
        block_end = title_columns[index + 1] - 1 if index + 1 < len(title_columns) else end_col
        block = _build_article_import_block(row, header_index, block_start, block_end)
        if block:
            blocks.append(block)
    return blocks


def _detect_article_import_table_blocks(rows: list[list[str]]) -> list[dict[str, Any]]:
    candidate_blocks: list[dict[str, Any]] = []
    for header_index, row in enumerate(rows):
        blocks: list[dict[str, Any]] = []
        for start_col, end_col in _article_import_header_segments(row):
            blocks.extend(_split_article_import_segment_by_titles(row, header_index, start_col, end_col))
        for block in blocks:
            if (
                int(block.get("score") or 0) >= 6
                and "title" in dict(block.get("field_map") or {})
            ):
                candidate_blocks.append(block)

    if not candidate_blocks:
        return []

    def overlaps(left: dict[str, Any], right: dict[str, Any]) -> bool:
        left_start = int(left.get("start_col") or 0)
        left_end = int(left.get("end_col") or left_start)
        right_start = int(right.get("start_col") or 0)
        right_end = int(right.get("end_col") or right_start)
        return left_start <= right_end and right_start <= left_end

    candidate_blocks.sort(key=lambda item: (
        int(item.get("header_index") or 0),
        int(item.get("start_col") or 0),
        -int(item.get("score") or 0),
    ))
    deduped: list[dict[str, Any]] = []
    seen_signatures: set[tuple[int, int, int, tuple[tuple[str, int], ...]]] = set()
    for block in candidate_blocks:
        signature = (
            int(block.get("header_index") or 0),
            int(block.get("start_col") or 0),
            int(block.get("end_col") or 0),
            tuple(sorted((str(field), int(index)) for field, index in dict(block.get("field_map") or {}).items())),
        )
        if signature in seen_signatures:
            continue
        seen_signatures.add(signature)
        deduped.append(block)

    for block in deduped:
        header_index = int(block.get("header_index") or 0)
        end_row = len(rows)
        for candidate in deduped:
            candidate_header = int(candidate.get("header_index") or 0)
            if candidate_header <= header_index:
                continue
            if overlaps(block, candidate):
                end_row = min(end_row, candidate_header)
        block["end_row"] = end_row
    return deduped


def _row_looks_like_import_header(row: list[str]) -> bool:
    fields = {_match_article_import_field(cell) for cell in row}
    fields.discard("")
    return "title" in fields and len(fields) >= 2


def _normalize_article_import_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = _stringify_table_cell(value)
    if not text:
        return ""
    if re.fullmatch(r"\d{13}", text):
        try:
            return datetime.fromtimestamp(int(text) / 1000).strftime("%Y-%m-%d")
        except Exception:
            return ""
    if re.fullmatch(r"\d{10}", text):
        try:
            return datetime.fromtimestamp(int(text)).strftime("%Y-%m-%d")
        except Exception:
            return ""
    parsed_date = parse_local_date(text)
    if parsed_date:
        return parsed_date.isoformat()
    normalized = text.replace("T", " ").replace("Z", " ")
    normalized = re.sub(r"([0-9]{4})\s*年\s*([0-9]{1,2})\s*月\s*([0-9]{1,2})\s*日?", r"\1-\2-\3", normalized)
    normalized = re.sub(r"([0-9]{4})\.([0-9]{1,2})\.([0-9]{1,2})", r"\1-\2-\3", normalized)
    normalized = normalized.replace("/", "-")
    compact_match = re.search(r"(?<!\d)((?:19|20)\d{2})(\d{2})(\d{2})(?!\d)", normalized)
    if compact_match:
        try:
            return datetime(
                int(compact_match.group(1)),
                int(compact_match.group(2)),
                int(compact_match.group(3)),
            ).strftime("%Y-%m-%d")
        except Exception:
            pass
    match = re.search(r"((?:19|20)\d{2}-\d{1,2}-\d{1,2})", normalized)
    if match:
        try:
            return datetime.strptime(match.group(1), "%Y-%m-%d").strftime("%Y-%m-%d")
        except Exception:
            pass
    return ""


def _normalize_article_import_media_type(value: Any, url: str, media_name: str) -> str:
    text = _stringify_table_cell(value).lower()
    if any(token in text for token in ("自媒体", "账号", "号", "博主", "self", "creator", "account")):
        return "selfmedia"
    if any(token in text for token in ("权威", "官媒", "官方", "媒体", "authority", "official")):
        return "authority"
    try:
        from core.article_store import classify_article_media_type

        return classify_article_media_type(url, media_name)
    except Exception:
        return "selfmedia"


_ARTICLE_IMPORT_PLATFORM_LABEL_ALIASES: dict[str, str] = {
    "头条": "今日头条",
    "今日头条": "今日头条",
    "头条号": "头条号",
    "搜狐": "搜狐",
    "搜狐号": "搜狐号",
    "知乎": "知乎",
    "博客园": "博客园",
    "公众号": "微信公众号",
    "微信": "微信公众号",
    "微信公众号": "微信公众号",
    "百家号": "百家号",
    "微博": "微博",
    "小红书": "小红书",
    "抖音": "抖音",
    "快手": "快手",
    "哔哩哔哩": "哔哩哔哩",
    "b站": "哔哩哔哩",
    "B站": "哔哩哔哩",
}


def _resolve_article_import_platform_label(value: Any) -> str:
    text = _stringify_table_cell(value)
    key = _normalize_import_header(text)
    if not key:
        return ""
    normalized_aliases = {
        _normalize_import_header(platform): label
        for platform, label in _ARTICLE_IMPORT_PLATFORM_LABEL_ALIASES.items()
    }
    return normalized_aliases.get(key, "")


def _split_article_import_platform_account(value: Any, *, allow_unknown_bracket_platform: bool = False) -> tuple[str, str]:
    text = _stringify_table_cell(value)
    if not text:
        return "", ""
    candidates: list[tuple[str, str, str]] = []
    bracket_match = re.fullmatch(r"(.+?)[（(【\[]([^（）()【】\[\]]+)[）)】\]]", text)
    if bracket_match:
        candidates.append(("bracket", bracket_match.group(1), bracket_match.group(2)))
    separator_match = re.fullmatch(r"(.+?)\s*(?:-|—|–|_|｜|\||/|／|:|：)\s*(.+)", text)
    if separator_match:
        candidates.append(("separator", separator_match.group(1), separator_match.group(2)))

    for kind, raw_platform, raw_account in candidates:
        platform = str(raw_platform or "").strip()
        account = str(raw_account or "").strip()
        normalized_platform = _resolve_article_import_platform_label(platform)
        if not normalized_platform and kind == "bracket" and allow_unknown_bracket_platform:
            normalized_platform = platform
        if normalized_platform and account and normalized_platform != account:
            return normalized_platform, account
    return "", ""


def _extract_article_import_items(file_name: str, data: bytes) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    tables = _read_article_import_tables(file_name, data)
    items: list[dict[str, Any]] = []
    mapping_details: list[dict[str, Any]] = []
    skipped_without_title = 0
    for sheet_name, rows in tables:
        blocks = _detect_article_import_table_blocks(rows)
        if not blocks:
            header_index, field_map, score = _detect_article_import_header(rows)
            blocks = [{
                "header_index": header_index,
                "field_map": field_map,
                "score": score,
                "start_col": 0,
                "end_col": len(rows[header_index]) - 1 if header_index >= 0 else 0,
            }]
        valid_blocks = [
            block for block in blocks
            if int(block.get("header_index", -1)) >= 0
            and int(block.get("score") or 0) >= 6
            and "title" in dict(block.get("field_map") or {})
        ]
        if not valid_blocks:
            continue
        for block_number, block in enumerate(valid_blocks, start=1):
            header_index = int(block.get("header_index") or 0)
            field_map = dict(block.get("field_map") or {})
            start_col = int(block.get("start_col") or 0)
            end_col = int(block.get("end_col") or max(field_map.values()))
            end_row = int(block.get("end_row") or len(rows))
            headers = rows[header_index]
            sheet_count = 0
            for row_index, row in enumerate(rows[header_index + 1:end_row], start=header_index + 1):
                segment = row[start_col:end_col + 1] if start_col < len(row) else []
                if not any(segment) or _row_looks_like_import_header(segment):
                    continue

                def value_for(field: str) -> str:
                    index = field_map.get(field)
                    if index is None or index >= len(row):
                        return ""
                    return _stringify_table_cell(row[index])

                title = value_for("title")
                if not title:
                    skipped_without_title += 1
                    continue
                supporting_values = [
                    value_for(field)
                    for field in ("url", "media_name", "published_at", "excerpt", "account_name")
                    if field in field_map
                ]
                if supporting_values and not any(supporting_values):
                    skipped_without_title += 1
                    continue
                item = {
                    "title": title,
                    "url": value_for("url"),
                    "media_name": value_for("media_name"),
                    "media_type": value_for("media_type") or str(block.get("default_media_type") or ""),
                    "published_at": value_for("published_at"),
                    "excerpt": value_for("excerpt"),
                    "account_name": value_for("account_name"),
                    "brand_name": value_for("brand_name"),
                    "task_name": value_for("task_name"),
                    "_sheet": sheet_name,
                    "_row": row_index + 1,
                }
                items.append(item)
                sheet_count += 1
            mapping_details.append({
                "sheet": sheet_name if len(valid_blocks) == 1 else f"{sheet_name}#{block_number}",
                "count": sheet_count,
                "columns": {
                    field: headers[index] if index < len(headers) else ""
                    for field, index in field_map.items()
                },
            })
    return items, {
        "sheets": mapping_details,
        "skipped_without_title": skipped_without_title,
    }
